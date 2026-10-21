"""Tests for cv_sync.py: the CV's publication and software lists, built from the spreadsheets.

These read the real spreadsheets and CV. Tests that write work on copies in a
temporary directory.
"""

import re
import shutil
from pathlib import Path

import openpyxl
import pytest

from build_publications import load_publications
from build_software import load_software
from citation_utils import format_paper_citation, paper_venue_detail, preprint_display_id
from cv_sync import (
    RoleSpan, author_key, begin_marker, default_role, end_marker, end_sentence, html_to_latex,
    latex_escape, load_role_spans, parse_role_overrides, replace_block, resolve_roles,
    sentence_fragment, sync_cv, trainee_key,
)

REPO = Path(__file__).parent.parent
CV = REPO / 'documents' / 'JRM_CV.tex'
PUBS = REPO / 'data' / 'publications.xlsx'
SOFTWARE = REPO / 'data' / 'software.xlsx'


def block(tex, name):
    return tex.split(begin_marker(name), 1)[1].split(end_marker(name), 1)[0]


@pytest.fixture
def workspace(tmp_path):
    """Copies of the CV and both spreadsheets."""
    for src in (CV, PUBS, SOFTWARE):
        shutil.copy(src, tmp_path / src.name)
    return tmp_path


def sync(ws, check=False):
    return sync_cv(ws / 'JRM_CV.tex', ws / 'publications.xlsx', ws / 'software.xlsx', check=check)


class TestCommittedCV:
    def test_cv_is_in_sync_with_the_spreadsheets(self):
        """The committed CV's lists are exactly what the spreadsheets generate.

        If this fails, run `python build.py` (or `python cv_sync.py`) from
        scripts/ and commit JRM_CV.tex with the spreadsheet change.
        """
        assert not sync_cv(CV, PUBS, SOFTWARE, check=True)

    def test_every_paper_and_chapter_is_listed_once(self):
        pubs = load_publications(PUBS)
        items = block(CV.read_text(encoding='utf-8'), 'PUBLICATIONS').count('\\item ')
        assert items == len(pubs['papers']) + len(pubs['chapters'])

    def test_every_software_package_is_listed_once(self):
        software = load_software(SOFTWARE)
        items = block(CV.read_text(encoding='utf-8'), 'SOFTWARE').count('\\item ')
        assert items == sum(len(rows) for rows in software.values())

    def test_lists_are_newest_first(self):
        tex = CV.read_text(encoding='utf-8')
        for name in ('PUBLICATIONS', 'SOFTWARE'):
            years = [int(y) for y in re.findall(r'^\\item .*?\((\d{4})\)', block(tex, name), re.M)]
            assert years == sorted(years, reverse=True), name

    def test_every_author_roles_cell_parses_and_names_real_authors(self):
        """build.py would fail on a bad override; this says which cell."""
        for path, loader in ((PUBS, load_publications), (SOFTWARE, load_software)):
            for sheet, rows in loader(path).items():
                for row in rows:
                    overrides = parse_role_overrides(row.get('author_roles', ''))
                    authors = [a.strip() for a in row.get('authors', '').split(',')]
                    missing = set(overrides) - set(authors)
                    assert not missing, (sheet, row.get('title') or row.get('name'), missing)


class TestSync:
    def test_rewrites_a_stale_block_and_then_is_idempotent(self, workspace):
        tex_path = workspace / 'JRM_CV.tex'
        tex = tex_path.read_text(encoding='utf-8')
        tex_path.write_text(replace_block(tex, 'SOFTWARE', '\\item stale'), encoding='utf-8')

        assert sync(workspace, check=True)
        assert '\\item stale' in tex_path.read_text(encoding='utf-8')  # check writes nothing
        assert sync(workspace)
        assert tex_path.read_text(encoding='utf-8') == tex
        assert not sync(workspace)

    def test_a_spreadsheet_edit_reaches_the_cv(self, workspace):
        wb = openpyxl.load_workbook(workspace / 'publications.xlsx')
        ws = wb['papers']
        hdr = [c.value for c in ws[1]]
        ws.cell(2, hdr.index('title') + 1).value = 'A brand-new title & more'
        ws.cell(2, hdr.index('cv_note') + 1).value = '(Under review)'
        wb.save(workspace / 'publications.xlsx')

        assert sync(workspace)
        tex = (workspace / 'JRM_CV.tex').read_text(encoding='utf-8')
        assert 'A brand-new title \\& more.' in block(tex, 'PUBLICATIONS')
        assert '(Under review)' in block(tex, 'PUBLICATIONS')

    def test_a_bad_override_names_the_item(self, workspace):
        wb = openpyxl.load_workbook(workspace / 'software.xlsx')
        ws = wb['python']
        hdr = [c.value for c in ws[1]]
        name = ws.cell(2, hdr.index('name') + 1).value
        ws.cell(2, hdr.index('author_roles') + 1).value = 'Nobody Q: grad'
        wb.save(workspace / 'software.xlsx')

        with pytest.raises(ValueError, match=re.escape(name)):
            sync(workspace)

    def test_text_outside_the_markers_is_untouched(self, workspace):
        tex_path = workspace / 'JRM_CV.tex'
        before = tex_path.read_text(encoding='utf-8')
        tex_path.write_text(replace_block(before, 'PUBLICATIONS', ''), encoding='utf-8')
        sync(workspace)
        after = tex_path.read_text(encoding='utf-8')
        outside = lambda t: re.sub(r'% BEGIN GENERATED.*?% END GENERATED \w+', '', t, flags=re.S)  # noqa: E731
        assert outside(after) == outside(before)


class TestReplaceBlock:
    def test_missing_markers_are_an_error(self):
        with pytest.raises(ValueError, match='found 0'):
            replace_block('no markers here', 'SOFTWARE', 'x')

    def test_duplicate_markers_are_an_error(self):
        pair = f"{begin_marker('SOFTWARE')}\n{end_marker('SOFTWARE')}\n"
        with pytest.raises(ValueError, match='found 2'):
            replace_block(pair + pair, 'SOFTWARE', 'x')

    def test_body_with_backslashes_is_inserted_verbatim(self):
        pair = f"{begin_marker('SOFTWARE')}\n{end_marker('SOFTWARE')}"
        out = replace_block(pair, 'SOFTWARE', '\\item \\textbf{x} \\1')
        assert '\\item \\textbf{x} \\1' in out


class TestRoles:
    def test_author_and_trainee_keys_agree(self):
        assert author_key('Fitzpatrick PC') == trainee_key('Paxton Fitzpatrick')
        assert author_key('Chen HT') == trainee_key('Hung-Tu Chen')
        assert author_key('van der Meer MAA') == ('van der meer', 'm')

    def test_mentorship_spans_load_from_the_real_cv(self):
        spans = load_role_spans(CV)
        assert RoleSpan('postdoc', 2016, 2018) in spans[('heusser', 'a')]
        roles = {s.role for s in spans[('fitzpatrick', 'p')]}
        assert {'undergrad', 'grad'} <= roles

    def test_overrides_parse(self):
        assert parse_role_overrides('Jolly E: postdoc; Zhu Z: none;') == {
            'Jolly E': 'postdoc', 'Zhu Z': None}
        assert parse_role_overrides('') == {}

    @pytest.mark.parametrize('cell', ['Jolly E postdoc', 'Jolly E: professor'])
    def test_bad_overrides_are_errors(self, cell):
        with pytest.raises(ValueError):
            parse_role_overrides(cell)

    def test_override_beats_default_and_director_is_never_marked(self):
        spans = {('heusser', 'a'): [RoleSpan('postdoc', 2016, 2018)]}
        roles = resolve_roles(['Heusser AC', 'Manning JR', 'Stranger X'], 2017,
                              {'Heusser AC': None}, spans)
        assert roles == [None, None, None]

    def test_override_for_someone_not_an_author_is_an_error(self):
        with pytest.raises(ValueError, match='Typo Q'):
            resolve_roles(['Heusser AC'], 2017, {'Typo Q': 'grad'}, {})

    def test_default_role_returns_a_known_role_or_none(self):
        spans = load_role_spans(CV)
        for person in spans.values():
            for year in (2010, 2017, 2021, 2026):
                assert default_role(person, year) in ('undergrad', 'grad', 'postdoc', None)


class TestLatex:
    def test_escape(self):
        assert latex_escape('50% & $5 #1 a_b') == '50\\% \\& \\$5 \\#1 a\\_b'
        assert latex_escape('Davos: a "smuggler"') == "Davos: a ``smuggler''"
        assert latex_escape('1—19') == '1--19'

    def test_html_description_to_latex(self):
        assert html_to_latex('Slice along the n<sup>th</sup> dimension') == \
            'Slice along the n\\textsuperscript{th} dimension'
        assert html_to_latex('With <a href="https://x.org/a_b" target="_blank">X & Y</a>.') == \
            'With \\href{https://x.org/a_b}{X \\& Y}.'
        assert html_to_latex('See [docs](https://d.io)') == 'See \\href{https://d.io}{docs}'

    def test_end_sentence(self):
        assert end_sentence('Plot data') == 'Plot data.'
        assert end_sentence('Is it ready?') == 'Is it ready?'
        assert end_sentence('Plot data. (Way better!)') == 'Plot data. (Way better!)'
        assert end_sentence("Like ``this.''") == "Like ``this.''"

    def test_sentence_fragment_keeps_names(self):
        assert sentence_fragment('Manage dependencies.') == 'manage dependencies.'
        assert sentence_fragment('LLM demos.') == 'LLM demos.'
        assert sentence_fragment('BrainIAK is great.') == 'BrainIAK is great.'
        assert sentence_fragment('Clustrix is a package.', 'Clustrix') == 'Clustrix is a package.'
        assert sentence_fragment('Android and web app.', 'Dream Stream') == 'Android and web app.'


class TestVenueDetail:
    @pytest.mark.parametrize('journal,pages,expected', [
        ('PsyArXiv', '10.31234/osf.io/mhxtd_v1', 'mhxtd'),
        ('bioRxiv', '10.1101/2021.02.22.432360v1', '2021.02.22.432360'),
        ('bioRxiv', '2021.05.10.443308v1', '2021.05.10.443308'),
        ('arXiv', '2510.21958', '2510.21958'),
        ('Zenodo', '10.5281/zenodo.3937849', '10.5281/zenodo.3937849'),
    ])
    def test_preprint_ids(self, journal, pages, expected):
        assert preprint_display_id(journal, pages) == expected

    def test_detail_forms(self):
        assert paper_venue_detail('Nature Communications', '17', '', '2055') == ', 17: 2055'
        assert paper_venue_detail('Cerebral Cortex', '30', '10', '5333—5345') == ', 30(10): 5333—5345'
        assert paper_venue_detail('Communications Biology', status='In press') == ': In press'
        assert paper_venue_detail('PsyArXiv', pages='10.31234/osf.io/bw9r2') == ': bw9r2'
        assert paper_venue_detail('Some Journal') == ''

    def test_website_cites_preprints_by_id(self):
        html = format_paper_citation('Ziman K, Manning JR', '2021', 'T', 'bioRxiv',
                                     pages='10.1101/2021.02.22.432360v1')
        assert html.endswith('<em>bioRxiv</em>: 2021.02.22.432360.')
