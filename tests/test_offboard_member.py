"""End-to-end tests for offboard_member.py.

Each test runs the real script on a sandboxed copy of the repository, so the
real spreadsheet, templates, CV source and build scripts are exercised and the
checked-in files are never touched.
"""

import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).parent.parent


@pytest.fixture
def sandbox(tmp_path):
    """A copy of the repo with no lab-manual submodule, so nothing is pushed."""
    for name in ("scripts", "templates", "documents", "css"):
        shutil.copytree(
            REPO / name, tmp_path / name,
            ignore=shutil.ignore_patterns("__pycache__", "*.aux", "*.log"),
        )
    (tmp_path / "data").mkdir()
    for f in (REPO / "data").iterdir():
        # people.xlsx is edited by the script; everything else (fonts, the
        # other sheets) is read-only, so a symlink is enough.
        if f.name == "people.xlsx":
            shutil.copy2(f, tmp_path / "data" / f.name)
        else:
            os.symlink(f, tmp_path / "data" / f.name)
    shutil.copy2(REPO / "people.html", tmp_path / "people.html")
    return tmp_path


def first_open_undergrad(root):
    """Name, CV display name, CV start year and photo of a current undergrad
    with an open CV entry and a photo on the people page."""
    wb = openpyxl.load_workbook(root / "data" / "people.xlsx")
    sheet = wb["members"]
    headers = [c.value for c in sheet[1]]
    cv = (root / "documents" / "JRM_CV.tex").read_text(encoding="utf-8")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if rec.get("role") != "undergrad" or not rec.get("name") or not rec.get("image"):
            continue
        display = rec["name"].title()
        m = re.search(r"\\item\s+" + re.escape(display) + r"\*?\s*\((\d{4})\s*--\s*\)", cv)
        if m:
            return rec["name"], display, m.group(1), rec["image"]
    pytest.fail("No current undergrad with an open CV entry in the real data")


def run_offboard(root, *args):
    return subprocess.run(
        [sys.executable, "offboard_member.py", *args],
        cwd=root / "scripts", capture_output=True, text=True, timeout=900,
    )


class TestOffboardRebuilds:
    def test_people_page_is_rebuilt(self, sandbox):
        name, _, _, image = first_open_undergrad(sandbox)
        before = (sandbox / "people.html").read_text(encoding="utf-8")
        assert image in before

        result = run_offboard(sandbox, name, "-y", "--end-year", "2099")
        out = result.stdout + result.stderr
        assert "people.html rebuilt successfully" in out, out

        after = (sandbox / "people.html").read_text(encoding="utf-8")
        assert after != before
        # Alumni are listed by name only, so the member's photo card is gone.
        assert image not in after

    def test_cv_outputs_are_rebuilt(self, sandbox):
        if shutil.which("xelatex") is None:
            pytest.skip("xelatex not available")
        name, display, start, _ = first_open_undergrad(sandbox)

        result = run_offboard(sandbox, name, "-y", "--end-year", "2099")
        out = result.stdout + result.stderr
        assert result.returncode == 0, out
        assert "CV rebuilt successfully" in out, out

        html = (sandbox / "documents" / "JRM_CV.html").read_text(encoding="utf-8")
        assert f"{display} ({start} – 2099)" in html

    def test_skip_rebuild_leaves_outputs_alone(self, sandbox):
        name, _, _, _ = first_open_undergrad(sandbox)
        before = (sandbox / "people.html").read_text(encoding="utf-8")

        result = run_offboard(sandbox, name, "-y", "--skip-rebuild")
        out = result.stdout + result.stderr
        assert result.returncode == 0, out
        assert "Rebuilding" not in out
        assert (sandbox / "people.html").read_text(encoding="utf-8") == before


class TestSpreadsheetAlumniOrder:
    def test_new_alumnus_row_keeps_the_sheet_in_order(self, sandbox):
        from people_order import order_key

        name, display, start, _ = first_open_undergrad(sandbox)
        result = run_offboard(sandbox, name, "-y", "--end-year", "2099", "--skip-rebuild")
        assert result.returncode == 0, result.stdout + result.stderr

        ws = openpyxl.load_workbook(sandbox / "data" / "people.xlsx")["alumni_undergrads"]
        rows = [(r[0], str(r[1])) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        keys = [order_key(n, re.search(r"\d{4}", y).group(0)) for n, y in rows]
        assert keys == sorted(keys)
        # Written with the CV's spelling of the name, not str.title() of the
        # spreadsheet's lowercase one.
        assert (display, f"{start}-2099") in rows

    def test_name_uses_the_cvs_capitalization(self, sandbox):
        # str.title() turns 'mcdonald' into 'Mcdonald'; the CV says McDonald.
        ws = openpyxl.load_workbook(sandbox / "data" / "people.xlsx")["members"]
        names = [r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
        if "miles mcdonald" not in names:
            pytest.skip("Miles McDonald is no longer a current member")
        result = run_offboard(sandbox, "miles mcdonald", "-y", "--end-year", "2099",
                              "--skip-rebuild")
        assert result.returncode == 0, result.stdout + result.stderr
        ws = openpyxl.load_workbook(sandbox / "data" / "people.xlsx")["alumni_undergrads"]
        alumni = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        assert "Miles McDonald" in alumni
        assert "Miles Mcdonald" not in alumni
